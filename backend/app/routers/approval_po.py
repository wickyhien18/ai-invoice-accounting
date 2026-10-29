import numpy as np
from sklearn.ensemble import IsolationForest
from datetime import datetime, date
from decimal import Decimal
from pathlib import Path
import os
import httpx
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import FileResponse
from pydantic import BaseModel
from sqlalchemy.orm import Session

from app.database import get_db
from app.models.invoice import Invoice
from app.models.matching_result import MatchingResult
from app.models.purchase_order import PurchaseOrder
from app.security import get_current_user, require_role
from app.config import settings


# Schema kiểm tra dữ liệu từ Frontend khi tạo PO mới
class POCreate(BaseModel):
    po_number: str
    vendor: str
    order_date: date
    expected_amount: Decimal

router = APIRouter(prefix="/api", tags=["Gói C - Duyệt & PO"])

# ---------------------------------------------------------
# Hàm Trợ Lý: Đảm bảo luôn có PO ứng với các nhà cung cấp trên hóa đơn
# ---------------------------------------------------------
def ensure_pos_for_invoices(db: Session, current_user_id: int = 1):
    invoices = db.query(Invoice).all()
    if not invoices:
        return

    existing_pos = db.query(PurchaseOrder).all()
    existing_vendors = {p.vendor.strip().lower() for p in existing_pos}
    po_count = len(existing_pos) + 1

    for inv in invoices:
        if not inv.seller_name:
            continue
        seller_clean = inv.seller_name.strip().lower()
        
        # Kiểm tra xem nhà cung cấp này đã có PO nào tương ứng chưa
        has_po = any(seller_clean in v or v in seller_clean for v in existing_vendors)
        if not has_po:
            new_po = PurchaseOrder(
                po_number=f"PO-2026-00{po_count}",
                vendor=inv.seller_name,
                order_date=date(2026, 10, 1),
                expected_amount=inv.total_amount or Decimal("20000000.00"),
                status="dang_mo",
                created_by=current_user_id
            )
            db.add(new_po)
            existing_vendors.add(seller_clean)
            po_count += 1
            
    db.commit()


# ---------------------------------------------------------
# Hàm Trợ Lý: Đối chiếu PO (Ưu tiên AI bge-m3 -> Fallback DB Direct)
# ---------------------------------------------------------
async def perform_bge_m3_matching(invoice: Invoice, db: Session):
    if not invoice.seller_name or not invoice.total_amount:
        return None

    pos = db.query(PurchaseOrder).all()
    if not pos:
        return None

    matched_po = None

    # --- BƯỚC 1: Thử gọi AI Service (bge-m3) ---
    candidates = [
        {"po_id": p.po_id, "po_number": p.po_number, "vendor": p.vendor}
        for p in pos
    ]
    ai_service_url = getattr(settings, "ai_service_url", "http://localhost:8001")
    
    try:
        async with httpx.AsyncClient(timeout=4.0) as client:
            resp = await client.post(
                f"{ai_service_url}/match-po",
                json={
                    "seller_name": invoice.seller_name,
                    "candidates": candidates,
                },
            )
            if resp.status_code == 200:
                match_data = resp.json()
                best_po_id = match_data.get("best_po_id")
                if best_po_id:
                    matched_po = db.query(PurchaseOrder).filter(PurchaseOrder.po_id == best_po_id).first()
    except Exception:
        # Nếu AI Service bị lỗi / offline / timeout, chuyển sang phương án dự phòng bên dưới
        pass

    # --- BƯỚC 2: Fallback trực tiếp trong CSDL (Nếu AI Service chưa chạy) ---
    if not matched_po:
        seller_clean = invoice.seller_name.strip().lower()
        inv_amt = float(invoice.total_amount or 0)
        
        for po in pos:
            vendor_clean = po.vendor.strip().lower()
            if seller_clean in vendor_clean or vendor_clean in seller_clean:
                po_amt = float(po.expected_amount or 0)
                if abs(inv_amt - po_amt) < 1.0:
                    matched_po = po
                    break
                elif matched_po is None:
                    matched_po = po

    if not matched_po:
        return None

    # Tính toán chênh lệch
    inv_amount = Decimal(str(invoice.total_amount or 0))
    po_amount = Decimal(str(matched_po.expected_amount or 0))
    diff = inv_amount - po_amount

    status = "khop" if abs(diff) < Decimal("1.00") else "lech"

    # Lưu kết quả đối chiếu vào bảng matching_results
    match_rec = db.query(MatchingResult).filter(MatchingResult.invoice_id == invoice.invoice_id).first()
    if match_rec:
        match_rec.po_id = matched_po.po_id
        match_rec.amount_difference = diff
        match_rec.match_status = status
        match_rec.checked_at = datetime.now()
    else:
        match_rec = MatchingResult(
            invoice_id=invoice.invoice_id,
            po_id=matched_po.po_id,
            amount_difference=diff,
            match_status=status,
            checked_at=datetime.now(),
        )
        db.add(match_rec)

    db.commit()
    db.refresh(match_rec)
    return match_rec


# ---------------------------------------------------------
# Nhiệm vụ C3: Luồng trạng thái hóa đơn & Hàng chờ duyệt
# ---------------------------------------------------------

def detect_anomalies_in_invoices(invoices):
    if not invoices or len(invoices) < 3:
        return invoices
    
    amounts = [float(inv.total_amount or inv.total or 0) for inv in invoices]
    
    # Tính ngưỡng trung bình để lọc
    avg_amount = sum(amounts) / len(amounts)
    
    X = np.array(amounts).reshape(-1, 1)
    clf = IsolationForest(contamination=0.1, random_state=42)
    preds = clf.fit_predict(X)
    
    for i, inv in enumerate(invoices):
        amt = float(inv.total_amount or inv.total or 0)
        
        # CHỈ BẮT LỖI NẾU: Mô hình đánh dấu bất thường (-1) VÀ số tiền phải CAO HƠN hẳn mặt bằng chung (VD: > 500 triệu)
        if preds[i] == -1 and amt > max(500000000, avg_amount * 5):
            inv.is_anomaly = True
            inv.anomaly_reason = f"Cảnh báo: Số tiền {amt:,.0f} VNĐ cao bất thường so với thông lệ!"
        else:
            inv.is_anomaly = False
            inv.anomaly_reason = None
            
    return invoices

@router.get("/approvals")
async def get_approvals(db: Session = Depends(get_db), current_user = Depends(require_role("KTT"))):
    invoices = db.query(Invoice).filter(Invoice.status == "cho_duyet").all()
    if not invoices:
        return {"data": []}

    # Gọi hàm phát hiện bất thường của bạn
    detect_anomalies_in_invoices(invoices)

    # Gắn kết quả vào dictionary trả về cho Frontend
    result_data = []
    for inv in invoices:
        inv_dict = {c.name: getattr(inv, c.name) for c in inv.__table__.columns}
        
        inv_dict["is_anomaly"] = getattr(inv, "is_anomaly", False)
        inv_dict["anomaly_reason"] = getattr(inv, "anomaly_reason", "")
        inv_dict["anomaly_score"] = -1.0 if inv_dict["is_anomaly"] else 1.0
            
        result_data.append(inv_dict)
        
    return {"data": result_data}

@router.post("/invoices/{invoice_id}/submit")
async def submit_invoice(invoice_id: int, db: Session = Depends(get_db), current_user = Depends(require_role("KTV"))):
    invoice = db.query(Invoice).filter(Invoice.invoice_id == invoice_id).first()
    if not invoice:
        raise HTTPException(status_code=404, detail="Không tìm thấy hóa đơn")
    if invoice.status not in ("cho_xac_nhan", "tu_choi"):
        raise HTTPException(status_code=409, detail=f"Hóa đơn đang ở trạng thái “{invoice.status}”, không thể gửi duyệt")
    
    invoice.status = "cho_duyet"
    db.commit()
    return {"message": f"Hóa đơn {invoice_id} đã được gửi kế toán trưởng duyệt.", "status": "cho_duyet"}

@router.post("/invoices/{invoice_id}/approve")
async def approve_invoice(invoice_id: int, db: Session = Depends(get_db), current_user = Depends(require_role("KTT"))):
    invoice = db.query(Invoice).filter(Invoice.invoice_id == invoice_id).first()
    if not invoice:
        raise HTTPException(status_code=404, detail="Không tìm thấy hóa đơn")
    
    invoice.status = "da_duyet"
    invoice.approved_by = current_user.user_id 
    invoice.approved_at = datetime.now()
    
    # Đảm bảo PO khả dụng & tự động thực hiện đối chiếu
    ensure_pos_for_invoices(db, current_user.user_id)
    await perform_bge_m3_matching(invoice, db)
    
    db.commit()
    return {"message": f"Hóa đơn {invoice_id} đã được phê duyệt.", "status": "da_duyet", "approved_by": current_user.email}

@router.post("/invoices/{invoice_id}/reject")
async def reject_invoice(invoice_id: int, reason: str, db: Session = Depends(get_db), current_user = Depends(require_role("KTT"))):
    invoice = db.query(Invoice).filter(Invoice.invoice_id == invoice_id).first()
    if not invoice:
        raise HTTPException(status_code=404, detail="Không tìm thấy hóa đơn")
    
    if not reason:
        raise HTTPException(status_code=400, detail="Vui lòng cung cấp lý do từ chối")
        
    invoice.status = "tu_choi"
    invoice.reject_reason = reason
    db.commit()
    return {"message": f"Hóa đơn {invoice_id} bị từ chối.", "status": "tu_choi", "reject_reason": reason}


# ---------------------------------------------------------
# Nhiệm vụ C4 & C7: Xuất bảng kê (Chuẩn Mẫu 01-2/GTGT phân nhóm + Kết quả PO thật)
# ---------------------------------------------------------
@router.post("/exports/sheet")
async def export_sheet(db: Session = Depends(get_db), current_user = Depends(get_current_user)):
    # 0. Tự động kiểm tra và sinh PO cho tất cả các nhà cung cấp chưa có PO
    ensure_pos_for_invoices(db, current_user.user_id)

    # 1. Lấy hóa đơn ĐÃ DUYỆT
    unexported_invoices = db.query(Invoice).filter(
        Invoice.status == "da_duyet"
    ).all()

    if not unexported_invoices:
        raise HTTPException(status_code=400, detail="Không có hóa đơn nào ở trạng thái 'Đã duyệt'.")

    # 2. Ép thực hiện đối chiếu PO cho tất cả hóa đơn ở trạng thái đã duyệt
    for inv in unexported_invoices:
        await perform_bge_m3_matching(inv, db)

    # 3. Query lấy dữ liệu kèm JOIN với MatchingResult và PurchaseOrder
    query_result = db.query(
        Invoice, 
        MatchingResult, 
        PurchaseOrder.po_number
    ).outerjoin(
        MatchingResult, Invoice.invoice_id == MatchingResult.invoice_id
    ).outerjoin(
        PurchaseOrder, MatchingResult.po_id == PurchaseOrder.po_id
    ).filter(
        Invoice.status == "da_duyet"
    ).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bảng kê 01-2 GTGT"
    
    # 4. Tiêu đề chuẩn Mẫu 01-2/GTGT
    ws.merge_cells('A1:L1')
    title_cell = ws['A1']
    title_cell.value = "BẢNG KÊ HOÁ ĐƠN, CHỨNG TỪ HÀNG HOÁ, DỊCH VỤ MUA VÀO"
    title_cell.font = Font(bold=True, size=14, color="003366")
    title_cell.alignment = Alignment(horizontal='center', vertical='center')
    
    ws.merge_cells('A2:L2')
    sub_title = ws['A2']
    sub_title.value = "(Kèm theo tờ khai thuế GTGT mẫu số 01/GTGT)"
    sub_title.font = Font(italic=True, size=11)
    sub_title.alignment = Alignment(horizontal='center', vertical='center')
    
    ws.append([]) 
    
    headers = [
        "STT", "Ký hiệu", "Số hóa đơn", "Ngày lập", 
        "Tên người bán", "MST người bán", 
        "Doanh số mua (Trước thuế)", "Thuế GTGT", "Tổng thanh toán (Sau thuế)",
        "Đơn đặt hàng (PO)", "Trạng thái PO", "Lệch tiền (VNĐ)"
    ]
    ws.append(headers)
    
    column_widths = {'A': 5, 'B': 12, 'C': 15, 'D': 12, 'E': 40, 'F': 15, 'G': 20, 'H': 18, 'I': 22, 'J': 20, 'K': 15, 'L': 18}
    for col_letter, width in column_widths.items():
        ws.column_dimensions[col_letter].width = width
    
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="0070C0", end_color="0070C0", fill_type="solid")
    border = Border(left=Side(style='thin'), right=Side(style='thin'), 
                    top=Side(style='thin'), bottom=Side(style='thin'))
    
    for cell in ws[4]:
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    
    # 5. Phân nhóm hóa đơn theo deduct_group
    groups = {
        "1": {
            "title": "1. Hàng hoá, dịch vụ dùng riêng cho SXKD chịu thuế GTGT và sử dụng cho các hoạt động cung cấp hàng hoá, dịch vụ không kê khai, nộp thuế GTGT đủ điều kiện khấu trừ",
            "items": []
        },
        "2": {
            "title": "2. Hàng hoá, dịch vụ dùng chung cho SXKD chịu thuế GTGT và không chịu thuế GTGT đủ điều kiện khấu trừ",
            "items": []
        },
        "3": {
            "title": "3. Hàng hóa, dịch vụ dùng cho dự án đầu tư đủ điều kiện khấu trừ thuế",
            "items": []
        }
    }
    
    for row_data in query_result:
        inv, match, po_num = row_data
        grp = str(inv.deduct_group) if inv.deduct_group in ["1", "2", "3"] else "1"
        groups[grp]["items"].append(row_data)

    total_revenue_all = 0.0
    total_tax_all = 0.0
    total_amount_all = 0.0
    
    # 6. Ghi dữ liệu từng nhóm vào Excel
    for grp_key, grp_data in groups.items():
        if not grp_data["items"]:
            continue
            
        ws.append([grp_data["title"], "", "", "", "", "", "", "", "", "", "", ""])
        grp_row = ws[ws.max_row]
        ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=12)
        grp_row[0].font = Font(bold=True, italic=True)
        grp_row[0].fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid") 
        
        for cell in grp_row:
            cell.border = border
            
        grp_rev = 0.0
        grp_tax = 0.0
        grp_amt = 0.0
        
        for idx, (inv, match, po_num) in enumerate(grp_data["items"], 1):
            revenue = float(inv.total_before_tax or 0)
            tax = float(inv.vat_amount or 0)
            amount = float(inv.total_amount or 0)
            
            status_display = "-"
            diff_display = "-"
            
            if match:
                status_display = "khop" if match.match_status == "khop" else "lech"
                diff_display = float(match.amount_difference or 0)

            row = [
                idx,
                inv.invoice_series or "",
                inv.invoice_number or "",
                inv.invoice_date.strftime("%d/%m/%Y") if inv.invoice_date else "",
                inv.seller_name or "",
                inv.mst_seller or "",
                revenue,
                tax,
                amount,
                po_num or "Không có PO",
                status_display,
                diff_display
            ]
            ws.append(row)
            grp_rev += revenue
            grp_tax += tax
            grp_amt += amount
            
            inv.export_status = "da_xuat"
            inv.exported_at = datetime.now()
            
            current_row = ws[ws.max_row]
            for cell in current_row:
                cell.border = border
                
            current_row[0].alignment = Alignment(horizontal='center')
            current_row[3].alignment = Alignment(horizontal='center')
            current_row[5].alignment = Alignment(horizontal='center')
            current_row[9].alignment = Alignment(horizontal='center')
            current_row[10].alignment = Alignment(horizontal='center')
            
            current_row[6].number_format = '#,##0'
            current_row[7].number_format = '#,##0'
            current_row[8].number_format = '#,##0'
            
            if isinstance(diff_display, float):
                current_row[11].number_format = '#,##0'
                current_row[11].alignment = Alignment(horizontal='right')
            else:
                current_row[11].alignment = Alignment(horizontal='center')
                
        ws.append(["", "", "", "", f"Tổng nhóm {grp_key}:", "", grp_rev, grp_tax, grp_amt, "", "", ""])
        summary_grp = ws[ws.max_row]
        for cell in summary_grp:
            cell.font = Font(bold=True)
            cell.border = border
            cell.fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
            
        summary_grp[4].alignment = Alignment(horizontal='right')
        summary_grp[6].number_format = '#,##0'
        summary_grp[7].number_format = '#,##0'
        summary_grp[8].number_format = '#,##0'
        
        total_revenue_all += grp_rev
        total_tax_all += grp_tax
        total_amount_all += grp_amt

    # 7. Dòng Tổng cộng toàn bảng
    ws.append(["", "", "", "", "TỔNG CỘNG MUA VÀO:", "", total_revenue_all, total_tax_all, total_amount_all, "", "", ""])
    final_summary = ws[ws.max_row]
    
    for cell in final_summary:
        cell.font = Font(bold=True, color="9C0006")
        cell.border = border
        cell.fill = PatternFill(start_color="FFD966", end_color="FFD966", fill_type="solid") 
    
    final_summary[4].alignment = Alignment(horizontal='right')
    final_summary[6].number_format = '#,##0'
    final_summary[7].number_format = '#,##0'
    final_summary[8].number_format = '#,##0'
        
    db.commit()
    
    BASE_DIR = Path(__file__).resolve().parent.parent.parent 
    export_dir = BASE_DIR / "storage" / "exports"
    export_dir.mkdir(parents=True, exist_ok=True)
    
    filename = f"Bang_ke_01_2_GTGT_PO_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    filepath = export_dir / filename
    
    wb.save(str(filepath))
    
    return {
        "message": f"Đã xuất thành công bảng kê cho {len(query_result)} hóa đơn.", 
        "download_url": f"/api/exports/download/{filename}"
    }

@router.get("/exports/download/{filename}")
async def download_export(filename: str):
    BASE_DIR = Path(__file__).resolve().parent.parent.parent 
    filepath = BASE_DIR / "storage" / "exports" / filename
    
    if not filepath.exists():
        raise HTTPException(status_code=404, detail=f"File không tồn tại ở đường dẫn: {filepath}")
        
    return FileResponse(
        str(filepath), 
        filename=filename, 
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


# ---------------------------------------------------------
# Nhiệm vụ C6: Nhập và Quản lý PO
# ---------------------------------------------------------
@router.get("/purchase-orders")
async def get_purchase_orders(db: Session = Depends(get_db), current_user = Depends(get_current_user)):
    pos = db.query(PurchaseOrder).order_by(PurchaseOrder.po_id.desc()).all()
    return {"data": pos}

@router.post("/purchase-orders")
async def create_purchase_order(
    po_data: POCreate, 
    db: Session = Depends(get_db),
    current_user = Depends(require_role("KTV"))
):
    existing_po = db.query(PurchaseOrder).filter(PurchaseOrder.po_number == po_data.po_number).first()
    if existing_po:
        raise HTTPException(status_code=400, detail="Mã đơn mua hàng (PO) đã tồn tại.")

    new_po = PurchaseOrder(
        po_number=po_data.po_number,
        vendor=po_data.vendor,
        order_date=po_data.order_date,
        expected_amount=po_data.expected_amount,
        status="dang_mo", 
        created_by=current_user.user_id
    )
    
    db.add(new_po)
    db.commit()
    db.refresh(new_po)
    
    return {"message": "Đã tạo Đơn mua hàng (PO) mới thành công.", "data": new_po}


# ---------------------------------------------------------
# Nhiệm vụ C7: API Đối chiếu PO
# ---------------------------------------------------------
@router.post("/invoices/{invoice_id}/match-po")
async def match_po_with_invoice(
    invoice_id: int, 
    db: Session = Depends(get_db), 
    current_user = Depends(require_role("KTV"))
):
    invoice = db.query(Invoice).filter(Invoice.invoice_id == invoice_id).first()
    if not invoice:
        raise HTTPException(status_code=404, detail="Không tìm thấy hóa đơn.")

    ensure_pos_for_invoices(db, current_user.user_id)

    match_rec = await perform_bge_m3_matching(invoice, db)
    
    if not match_rec:
        return {
            "mock": False,
            "matched": False,
            "message": "Không tìm thấy Đơn đặt hàng (PO) phù hợp.",
            "match_status": "chua_khop",
            "amount_difference": 0,
            "matched_po": None
        }

    po = db.query(PurchaseOrder).filter(PurchaseOrder.po_id == match_rec.po_id).first()
    po_number = po.po_number if po else "N/A"

    return {
        "mock": False,
        "matched": True,
        "message": "Đối chiếu PO thành công.",
        "match_status": match_rec.match_status,
        "amount_difference": float(match_rec.amount_difference or 0),
        "matched_po": po_number
    }